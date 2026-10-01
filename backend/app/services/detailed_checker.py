from __future__ import annotations

import re
from typing import Any, Dict, List, Optional

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    Document = None
    WD_ALIGN_PARAGRAPH = None
    DOCX_AVAILABLE = False

from app.models.document_model import (
    DocumentModel,
    FigureModel,
    ParagraphModel,
    ReferenceModel,
    SectionModel,
    TableModel,
)


class DocxParser:
    def __init__(self):
        self.docx_available = DOCX_AVAILABLE

    def parse(self, payload: Dict[str, Any]) -> DocumentModel:
        doc = DocumentModel()

        doc.title = str(payload.get("title", "")).strip()
        doc.english_title = str(payload.get("english_title", "")).strip()
        doc.abstract = str(payload.get("abstract", "")).strip()
        doc.keywords = [
            str(item).strip() for item in payload.get("keywords", []) if str(item).strip()
        ]

        for section in payload.get("sections", []) or []:
            doc.sections.append(
                SectionModel(
                    title=str(section.get("title", "")).strip(),
                    level=int(section.get("level", 1) or 1),
                    start_page=section.get("start_page"),
                    end_page=section.get("end_page"),
                )
            )

        for paragraph in payload.get("paragraphs", []) or []:
            doc.paragraphs.append(
                ParagraphModel(
                    text=str(paragraph.get("text", "")).strip(),
                    font=str(paragraph.get("font", "")).strip(),
                    font_size=str(paragraph.get("font_size", "")).strip(),
                    align=str(paragraph.get("align", "left")).strip() or "left",
                    line_spacing=str(paragraph.get("line_spacing", "1.25")).strip(),
                    section_index=paragraph.get("section_index"),
                )
            )

        for figure in payload.get("figures", []) or []:
            doc.figures.append(
                FigureModel(
                    caption=str(figure.get("caption", "")).strip(),
                    alt_text=str(figure.get("alt_text", "")).strip(),
                    page=figure.get("page"),
                )
            )

        for table in payload.get("tables", []) or []:
            doc.tables.append(
                TableModel(
                    caption=str(table.get("caption", "")).strip(),
                    page=table.get("page"),
                )
            )

        for ref in payload.get("references", []) or []:
            if isinstance(ref, dict):
                doc.references.append(
                    ReferenceModel(
                        id=int(ref.get("id", 0) or 0),
                        text=str(ref.get("text", "")).strip(),
                        ref_type=str(ref.get("type", "journal")).strip(),
                    )
                )

        doc.pages = payload.get("pages", {}) or {"front_pages": "roman", "main_pages": "arabic"}
        return doc

    def extract_from_docx_bytes(self, file_bytes: bytes) -> DocumentModel:
        if not self.docx_available:
            raise ImportError("python-docx is not installed. Please install it with: pip install python-docx")

        from io import BytesIO

        document = Document(BytesIO(file_bytes))
        doc = DocumentModel()
        word_paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            font_name = paragraph.runs[0].font.name if paragraph.runs else "宋体"
            font_size = str(paragraph.runs[0].font.size) if paragraph.runs and paragraph.runs[0].font.size else "12pt"
            align = self._get_alignment(paragraph)

            if self._is_heading(paragraph):
                level = self._heading_level(paragraph)
                doc.sections.append(SectionModel(title=text, level=level, start_page=1))
                continue

            if self._is_title(paragraph):
                if not doc.title:
                    doc.title = text
                continue

            if "摘要" in text and len(text) < 200:
                if not doc.abstract:
                    doc.abstract = text.replace("摘要", "").strip()
                continue

            if "关键词" in text:
                keywords_text = text.replace("关键词", "").replace("Keywords", "").strip()
                doc.keywords = [k.strip() for k in re.split(r"[,，]", keywords_text) if k.strip()]
                continue

            word_paragraphs.append(
                ParagraphModel(
                    text=text,
                    font=font_name,
                    font_size=font_size,
                    align=align,
                    line_spacing="1.25",
                    section_index=len(doc.sections),
                )
            )

        doc.paragraphs = word_paragraphs

        for table in document.tables:
            values = [cell.text.strip() for row in table.rows for cell in row.cells if cell.text.strip()]
            caption = " ".join(values[:5]).strip() or ""
            doc.tables.append(TableModel(caption=caption, page=1))

        for rel in document.part.rels.values():
            if "image" in str(rel.target_ref).lower():
                doc.figures.append(FigureModel(caption="图", page=1))

        doc.pages = {"front_pages": "roman", "main_pages": "arabic"}
        return doc

    def _is_heading(self, paragraph) -> bool:
        if paragraph.style and "Heading" in paragraph.style.name:
            return True
        if paragraph.runs and paragraph.runs[0].font.bold and len(paragraph.text.strip()) < 60:
            return True
        return False

    def _heading_level(self, paragraph) -> int:
        style_name = getattr(getattr(paragraph, "style", None), "name", "")
        if style_name == "Heading 1":
            return 1
        if style_name == "Heading 2":
            return 2
        if style_name == "Heading 3":
            return 3
        return 1

    def _is_title(self, paragraph) -> bool:
        text = paragraph.text.strip()
        if len(text) > 100:
            return False
        if paragraph.runs and paragraph.runs[0].font.bold:
            return True
        return False

    def _get_alignment(self, paragraph) -> str:
        if WD_ALIGN_PARAGRAPH is None:
            return "left"
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            return "center"
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            return "right"
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
            return "justify"
        return "left"
